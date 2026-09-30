"""Inspect the EBO pricing models: what did a real buyer assume and pay?

These are the cash-flow models an actual purchaser of Ginnie Mae early buyouts used to
bid on pools, by counterparty, 2018 to February 2020. If they contain bid prices, yields
and repool assumptions, they replace the assumed premium in the paper's valuation with
something observed.
"""
import os, glob, warnings
warnings.filterwarnings("ignore")
import openpyxl

D = r"F:\Ekhoni_Lagbe_Na\Kristine\Propensity Modelling\CF outputs"
files = sorted(glob.glob(os.path.join(D, "*.xls*")), key=os.path.getsize, reverse=True)
print(f"{len(files)} pricing files\n")
for f in files:
    print(f"  {os.path.getsize(f)/1e6:7.1f} MB  {os.path.basename(f)}")

KEY = ["price", "bid", "yield", "irr", "repool", "cure", "advance", "upb", "wac",
       "coupon", "severity", "loss", "discount", "purchase", "px", "dollar"]

TARGETS = [f for f in files if any(k in os.path.basename(f).lower()
           for k in ["wells", "chase", "citi", "bana", "ditech", "ebo1014"])][:6]
for f in TARGETS:
    print("\n" + "=" * 88)
    print(os.path.basename(f))
    print("=" * 88, flush=True)
    try:
        wb = openpyxl.load_workbook(f, data_only=True, read_only=True)
    except Exception as e:
        print("   could not open:", type(e).__name__, e); continue
    print("   sheets:", wb.sheetnames[:20])
    for sn in wb.sheetnames[:8]:
        ws = wb[sn]
        hits = []
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 1, 90),
                                max_col=min(ws.max_column or 1, 14)):
            for c in row:
                v = c.value
                if isinstance(v, str) and 2 < len(v) < 46:
                    lv = v.lower()
                    if any(k in lv for k in KEY):
                        nb = [x.value for x in row if isinstance(x.value, (int, float))]
                        if nb:
                            hits.append((c.coordinate, v.strip(), nb[:4]))
        if hits:
            print(f"\n   --- sheet '{sn}' ({ws.max_row}x{ws.max_column}) ---")
            for co, lab, nb in hits[:14]:
                nbs = ", ".join(f"{x:,.4g}" for x in nb)
                print(f"      {co:>6}  {lab:<44} {nbs}")
    wb.close()
