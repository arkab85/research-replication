"""Build anonymous editable manuscripts and an author-only submission note."""
from pathlib import Path
import re,subprocess
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
R=Path(__file__).resolve().parent

def format_doc(path,kind):
 d=Document(path);sec=d.sections[0]
 sec.page_width=Inches(8.5);sec.page_height=Inches(11)
 sec.top_margin=Inches(.8);sec.bottom_margin=Inches(.8);sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
 for name in ['Normal','Body Text','First Paragraph']:
  st=d.styles[name];st.font.name='Times New Roman';st.font.size=Pt(11 if kind=='supp' else 11.5);st.font.color.rgb=RGBColor(0,0,0)
  st.paragraph_format.line_spacing=1.18 if kind=='supp' else 1.28
  st.paragraph_format.space_after=Pt(6);st.paragraph_format.widow_control=True
 for name,size in [('Title',19),('Subtitle',13),('Heading 1',13),('Heading 2',11.5)]:
  st=next(z for z in d.styles if z.style_id==name.replace(' ',''));st.font.name='Times New Roman';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st.font.bold=name!='Subtitle'
  st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(6);st.paragraph_format.keep_with_next=True
 for st in d.styles:
  for rf in st.element.findall('.//'+qn('w:rFonts')):
   for a in list(rf.attrib):
    if 'theme' in a.lower():del rf.attrib[a]
  for c in st.element.findall('.//'+qn('w:color')):
   for a in ['themeColor','themeTint','themeShade']:c.attrib.pop(qn('w:'+a),None)
 for p in d.paragraphs:
  if p.style.name in ['Title','Subtitle']:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  if p.text.startswith(('Table ','Fig. ')):
   p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(8)
  if p._p.xpath('.//w:drawing'):
   p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
  if p.text.startswith('Notes:'):
   p.paragraph_format.line_spacing=1.05
   for run in p.runs:run.font.size=Pt(9.5)
  if p.text=='Cover letter draft':p.paragraph_format.page_break_before=True
  for el in p._p.findall('.//'+qn('w:pBdr')):el.getparent().remove(el)
 mainwidth=[[.6,.9,1.3,1.3,1.5],[.8,1.05,.55,1.1,1.65,1.55],[2.3,1.0,2.4,1.0],[2.9,.8,1.0,2.1],[1.0,.8,1.1,2.5,1.0],[2.7,.75,1.0,1.25,1.1]]
 for ti,t in enumerate(d.tables):
  n=len(t.columns)
  if kind=='main':width=mainwidth[ti] if ti < len(mainwidth) else [6.8/n]*n
  elif n==6:width=[2.1,.58,.55,.83,1.85,.65]
  elif n==5:width=[.85,1.35,1.35,1.35,1.35]
  elif n==4:width=[2.4,1.0,2.0,1.0] if ti==0 else [1.5,1.3,2.4,1.3]
  else:width=[6.8/n]*n
  width=[x*6.8/sum(width) for x in width]
  t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
  for col,w in zip(t.columns,width):col.width=Inches(w)
  for ri,row in enumerate(t.rows):
   trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
   if ri==0:trpr.append(OxmlElement('w:tblHeader'))
   for ci,cell in enumerate(row.cells):
    cell.width=Inches(width[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    pr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
    for edge in ['top','left','bottom','right']:
     el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    pr.append(borders);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EDF2' if ri==0 else 'FFFFFF');pr.append(sh)
    mar=OxmlElement('w:tcMar')
    for edge in ['top','bottom','left','right']:
     el=OxmlElement('w:'+edge);el.set(qn('w:w'),'65');el.set(qn('w:type'),'dxa');mar.append(el)
    pr.append(mar)
    for p in cell.paragraphs:
     p.paragraph_format.line_spacing=1.05;p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
     p.paragraph_format.keep_with_next=True if kind=='main' else (ri==0)
     p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
     table_size = 9 if kind=='supp' else (9 if ti==5 else 9.5)
     for run in p.runs:run.font.name='Times New Roman';run.font.size=Pt(table_size);run.font.bold=ri==0
 for i,p in enumerate(d.paragraphs[:-1]):
  nxt=d.paragraphs[i+1]
  if not p.text.strip() and not nxt.text.strip() and p._p.xpath('.//m:oMath') and nxt._p.xpath('.//m:oMath'):p.paragraph_format.keep_with_next=True
  if p.text.rstrip().endswith((' is',' are')) and nxt._p.xpath('.//m:oMath'):p.paragraph_format.keep_with_next=True
 inrefs=False
 for p in d.paragraphs:
  if p.text=='References':inrefs=True;continue
  if inrefs:
   p.paragraph_format.left_indent=Inches(.22);p.paragraph_format.first_line_indent=Inches(-.22);p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_after=Pt(3)
   for run in p.runs:run.font.size=Pt(10.5)
   for rp in p._p.xpath('.//w:r'):
    pr=rp.find(qn('w:rPr'))
    if pr is None:pr=OxmlElement('w:rPr');rp.insert(0,pr)
    sz=pr.find(qn('w:sz'))
    if sz is None:sz=OxmlElement('w:sz');pr.append(sz)
    sz.set(qn('w:val'),'21')
 for shape in d.inline_shapes:
  if shape.width>Inches(6.7):shape.height=int(shape.height*Inches(6.7)/shape.width);shape.width=Inches(6.7)
 p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)
 d.core_properties.author='';d.core_properties.last_modified_by='';d.core_properties.comments='';d.core_properties.subject=''
 for el in d.element.xpath('//*[@w:author]'):el.set(qn('w:author'),'')
 d.save(path)

for source,name,kind in [('manuscript.md','JUE_Spatial_Consumption.docx','main'),('supplement.md','JUE_Supplement.docx','supp'),('submission_note.md','JUE_Submission_Note.docx','note'),('highlights.md','JUE_Highlights.docx','note'),('title_page.md','JUE_Title_Page.docx','note')]:
 if not (R/source).exists():continue
 s=(R/source).read_text()
 s=re.sub(r'\\tag\{([^}]+)\}',lambda m:r'\qquad\text{('+m.group(1)+')}',s)
 (R/'document_input.md').write_text(s)
 subprocess.run(['pandoc','document_input.md','--standalone','--from=markdown-implicit_figures','-o',name],cwd=R,check=True)
 format_doc(R/name,kind)
 # Editable TeX source is provided for the two manuscripts.
 if kind!='note':
  subprocess.run(['pandoc',source,'--standalone','--from=markdown-implicit_figures','--pdf-engine=xelatex','-V','mainfont=DejaVu Serif','-V','geometry:margin=0.8in','-V','fontsize:11pt','-o','Main.tex' if kind=='main' else 'Supplement.tex'],cwd=R,check=True)
(R/'document_input.md').unlink()
print('Built manuscript, supplement, note, highlights, title page, and TeX sources.')
